from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Anexo Codigo Fuente Esencial EMS.docx"


FILES: list[tuple[str, str]] = [
    ("Configuracion del proyecto", "package.json"),
    ("Configuracion del proyecto", "vite.config.ts"),
    ("Configuracion del proyecto", "tsconfig.json"),
    ("Configuracion del proyecto", ".env.example"),
    ("Entrada frontend", "src/main.tsx"),
    ("Entrada frontend", "src/App.tsx"),
    ("Contrato TypeScript", "src/types/index.ts"),
    ("Cliente API", "src/services/apiClient.ts"),
    ("Servicios frontend", "src/services/dispatchEngine.ts"),
    ("Servicios frontend", "src/services/predictionEngine.ts"),
    ("Servicios frontend", "src/services/simulationEngine.ts"),
    ("Servicios frontend", "src/services/statsEngine.ts"),
    ("Componentes frontend", "src/components/Navbar.tsx"),
    ("Componentes frontend", "src/components/LiveKPIHeader.tsx"),
    ("Componentes frontend", "src/components/SimulationControls.tsx"),
    ("Componentes frontend", "src/components/PredictionPanel.tsx"),
    ("Componentes frontend", "src/components/ComparisonView.tsx"),
    ("Componentes frontend", "src/components/AssistantPanel.tsx"),
    ("Componentes frontend", "src/components/DigitalTwin3D.tsx"),
    ("Modulo historico frontend", "src/components/historical/HistoricalWorkspace.tsx"),
    ("Modulo historico frontend", "src/components/historical/HistoricalScene.tsx"),
    ("Modulo historico frontend", "src/components/historical/HistoricalMap.tsx"),
    ("Modulo historico frontend", "src/components/historical/ModelPerformancePanel.tsx"),
    ("Modulo historico frontend", "src/components/historical/CrispDmPanel.tsx"),
    ("Backend FastAPI", "backend/app/main.py"),
    ("Backend FastAPI", "backend/app/core/config.py"),
    ("Backend FastAPI", "backend/app/core/paths.py"),
    ("Rutas backend", "backend/app/api/health.py"),
    ("Rutas backend", "backend/app/api/catalog.py"),
    ("Rutas backend", "backend/app/api/datasets.py"),
    ("Rutas backend", "backend/app/api/simulations.py"),
    ("Rutas backend", "backend/app/api/predictions.py"),
    ("Rutas backend", "backend/app/api/comparisons.py"),
    ("Rutas backend", "backend/app/api/model_performance.py"),
    ("Rutas backend", "backend/app/api/methodology.py"),
    ("Rutas backend", "backend/app/api/assistant.py"),
    ("Servicios backend", "backend/app/services/historical_data_service.py"),
    ("Servicios backend", "backend/app/services/ems_simulator.py"),
    ("Servicios backend", "backend/app/services/simulation_service.py"),
    ("Servicios backend", "backend/app/services/comparison_service.py"),
    ("Servicios backend", "backend/app/services/model_performance_service.py"),
    ("Servicios backend", "backend/app/services/methodology_service.py"),
    ("Servicios backend", "backend/app/services/assistant_service.py"),
    ("Servicios backend", "backend/app/services/run_storage.py"),
    ("Pipeline investigativo", "research/pipeline/nyc_ems.py"),
    ("Experimentos investigativos", "research/experiments/audit_phase8c.py"),
    ("Experimentos investigativos", "research/experiments/fleet_sensitivity.py"),
    ("Visualizacion y reportes", "research/visualization/model_performance.py"),
    ("Visualizacion y reportes", "research/visualization/crisp_dm_methodology.py"),
    ("Visualizacion y reportes", "research/visualization/fleet_sensitivity.py"),
    ("Pruebas", "backend/tests/test_ems_simulator.py"),
    ("Pruebas", "research/tests/test_nyc_ems.py"),
    ("Pruebas", "research/tests/test_fleet_sensitivity.py"),
]


def read_text(relative: str) -> str:
    path = ROOT / relative
    return path.read_text(encoding="utf-8", errors="replace")


def line_count(relative: str) -> int:
    return len(read_text(relative).splitlines())


def configure_styles(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Calibri"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Calibri"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.name = "Calibri"
    styles["Heading 3"].font.size = Pt(11)

    if "CodeBlock" not in styles:
        style = styles.add_style("CodeBlock", 1)
    else:
        style = styles["CodeBlock"]
    style.font.name = "Consolas"
    style.font.size = Pt(7)


def add_code_block(document: Document, code: str) -> None:
    chunk_size = 12000
    for start in range(0, len(code), chunk_size):
        paragraph = document.add_paragraph(style="CodeBlock")
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.add_run(code[start : start + chunk_size])


def add_page_break(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_document() -> None:
    document = Document()
    configure_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    document.add_heading("Anexo de codigo fuente esencial", 0)
    document.add_paragraph(
        "Proyecto: Gemelo digital EMS con datos historicos NYC, simulacion, "
        "modelos predictivos, asistente analitico y reportes CRISP-DM."
    )
    document.add_paragraph(
        f"Documento generado automaticamente el {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}."
    )
    document.add_paragraph(
        "Criterio de inclusion: se incorporan los archivos que explican la "
        "arquitectura funcional del software. Se excluyen datasets, artefactos "
        "generados, dependencias, builds y archivos masivos no necesarios para "
        "evaluar la implementacion."
    )

    total_lines = sum(line_count(relative) for _, relative in FILES if (ROOT / relative).exists())
    document.add_paragraph(f"Archivos incluidos: {len(FILES)}. Lineas fuente aproximadas: {total_lines}.")

    document.add_heading("Indice de archivos incluidos", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Seccion"
    hdr[1].text = "Archivo"
    hdr[2].text = "Lineas"
    for section_name, relative in FILES:
        row = table.add_row().cells
        row[0].text = section_name
        row[1].text = relative
        row[2].text = str(line_count(relative)) if (ROOT / relative).exists() else "No encontrado"

    add_page_break(document)

    current_section = ""
    for section_name, relative in FILES:
        path = ROOT / relative
        if not path.exists():
            continue
        if section_name != current_section:
            document.add_heading(section_name, level=1)
            current_section = section_name
        document.add_heading(relative, level=2)
        document.add_paragraph(f"Ruta local: {relative}")
        add_code_block(document, read_text(relative))
        add_page_break(document)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
