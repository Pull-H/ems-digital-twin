from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SNAPSHOT = "20260909T040537085584Z-56d48fe7ce9b"
SENSITIVITY = ROOT / "reports" / "nyc_ems" / SNAPSHOT / "phase8d" / "583b13134c76f924dffd"
OUTPUT = ROOT / "docs" / "Guia de exposicion Gemelo Digital EMS.docx"

NAVY = "14283F"
TEAL = "087F8C"
PALE = "EAF4F5"
LIGHT = "F2F5F8"
GRAY = "506173"
ORANGE = "C15B27"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell_text(cell, text: str, *, bold=False, color=NAVY, size=9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF", size=8.5)
        shade(table.rows[0].cells[index], NAVY)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            set_cell_text(cells[index], value, size=8.5)
            if len(table.rows) % 2 == 0:
                shade(cells[index], LIGHT)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(document: Document, label: str, text: str, fill=PALE, accent=TEAL):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    title = paragraph.add_run(label + "\n")
    title.bold = True
    title.font.name = "Calibri"
    title.font.size = Pt(10)
    title.font.color.rgb = RGBColor.from_string(accent)
    body = paragraph.add_run(text)
    body.font.name = "Calibri"
    body.font.size = Pt(9.5)
    body.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(document: Document, text: str, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_number(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)
    return paragraph


def add_qa(document: Document, question: str, answer: str):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(question)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    body = document.add_paragraph(answer)
    body.paragraph_format.left_indent = Inches(0.18)
    body.paragraph_format.space_after = Pt(5)


def add_figure(document: Document, filename: str, caption: str, interpretation: str):
    path = SENSITIVITY / "figures" / filename
    if not path.exists():
        add_callout(document, "Grafica no disponible localmente", f"No se encontro {path.relative_to(ROOT)}.", fill="FFF3E8", accent=ORANGE)
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.35))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(3)
    run = cap.add_run(caption)
    run.bold = True
    run.font.size = Pt(9)
    add_callout(document, "Como leerla", interpretation)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Gemelo Digital EMS  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in (("Title", 30, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 11, NAVY)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    add_page_number(section.footer.paragraphs[0])


def build():
    document = Document()
    configure(document)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("GUÍA PARA EXPOSICIÓN")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    document.add_heading("Gemelo Digital para Servicios Médicos de Emergencia", 0)
    subtitle = document.add_paragraph("Datos históricos, predicción de demanda y simulación de despacho")
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    document.add_paragraph("Material de apoyo para exposición, demostración y preguntas del público")
    document.add_paragraph("Desarrollado en la asignatura Trabajo de Investigación, 10.º ciclo")
    document.add_paragraph(f"Versión de la guía: {date.today().isoformat()}")
    add_callout(
        document,
        "Idea central en una frase",
        "El sistema transforma registros históricos de llamadas EMS en pronósticos horarios de demanda y permite explorar, mediante simulación, cómo distintas reglas de despacho se comportarían bajo escenarios controlados.",
    )
    document.add_page_break()

    document.add_heading("1. Presentación rápida", level=1)
    document.add_heading("Elevator pitch (30 segundos)", level=2)
    document.add_paragraph(
        "Desarrollamos un prototipo de gemelo digital para estudiar la atención de emergencias médicas. Procesa datos públicos de incidentes de Nueva York, pronostica cuántas llamadas podrían ocurrir por hora en cada borough y utiliza un simulador para comparar políticas de despacho bajo diferentes niveles de flota. La herramienta permite visualizar demanda, unidades, colas y tiempos; sus resultados son experimentales y dependen de supuestos explícitos, por lo que no reemplaza un sistema operativo de despacho ni demuestra todavía mejoras en el mundo real."
    )
    document.add_heading("Exposición sugerida (3 a 5 minutos)", level=2)
    for item in [
        "Problema: los servicios de emergencia deben responder a una demanda que cambia según el lugar y la hora, mientras administran recursos limitados.",
        "Propuesta: integrar datos, pronóstico y simulación en una sola herramienta para analizar escenarios antes de tomar decisiones operativas.",
        "Datos: 405,998 registros públicos de NYC EMS entre abril y junio de 2026, obtenidos de la plataforma Socrata de la Ciudad de Nueva York.",
        "Método: limpieza y cohortes, panel por borough y hora, comparación de tres modelos de regresión y simulación de dos políticas de despacho.",
        "Hallazgo: los modelos reducen el error frente al baseline temporal en la prueba; la ventaja de una política de despacho no es uniforme y cambia con la disponibilidad supuesta de ambulancias.",
        "Aporte y límite: el prototipo ofrece una plataforma reproducible para explorar decisiones; para aplicarlo en Trujillo se requerirían datos locales, calibración y validación con los operadores del servicio.",
    ]:
        add_number(document, item)
    add_callout(document, "Frase para cerrar", "La contribución actual es una herramienta de análisis y experimentación. Su valor está en hacer visibles los datos, los supuestos y las consecuencias posibles de cada escenario, no en afirmar que ya optimiza un servicio real.", fill="FFF3E8", accent=ORANGE)

    document.add_heading("2. ¿Qué hace el sistema?", level=1)
    document.add_paragraph(
        "La aplicación combina una interfaz React/Vite, una API en Python/FastAPI, servicios de simulación y artefactos de investigación. El flujo toma un snapshot local del dataset, consulta sus perfiles y predicciones, configura una corrida y presenta el comportamiento del escenario en vistas 2D/3D, indicadores y tablas. Los scripts de investigación preparan datos y entrenan modelos; el backend carga los artefactos resultantes para la experiencia histórica."
    )
    add_table(document, ["Parte", "Qué hace", "Cómo explicarla"], [
        ["Frontend", "Controles, mapas, escena 3D, paneles y tablas.", "Es la interfaz que permite explorar y observar una corrida."],
        ["Backend FastAPI", "Expone endpoints, carga datos/modelos y controla el motor.", "Es el servicio Python que entrega los datos y ejecuta la simulación."],
        ["Pipeline research", "Descarga, valida, limpia, crea variables, entrena y produce reportes.", "Es el flujo reproducible que transforma datos crudos en artefactos analíticos."],
        ["Motor SimPy", "Procesa llegadas, cola, despacho, viaje, atención y liberación de unidades.", "Representa eventos y recursos a través del tiempo simulado."],
        ["Asistente analítico", "Consulta documentos, métricas y estado de corrida mediante un flujo LangGraph.", "Ayuda a interpretar evidencia del proyecto; no decide despachos."],
    ])
    document.add_heading("Qué significa aquí ‘gemelo digital’", level=2)
    document.add_paragraph(
        "Se usa como representación computacional de un proceso EMS por zonas, alimentada y calibrada con registros históricos y conectada a un simulador. No es un gemelo conectado en tiempo real a ambulancias, GPS, hospitales o centrales de despacho. La denominación describe el objetivo del prototipo de investigación, mientras que la sincronización operacional real permanece fuera del alcance actual."
    )
    add_callout(document, "Alcance que conviene decir en voz alta", "Es un gemelo digital retrospectivo y experimental, no un sistema de despacho en vivo ni una réplica completa de FDNY EMS.", fill="FFF3E8", accent=ORANGE)
    document.add_heading("Estado del proyecto", level=2)
    document.add_paragraph(
        "La base funcional ya integra extracción y procesamiento históricos, tres modelos de demanda, FastAPI, frontend, simulador, asistente analítico y reportes; además, las fases 8C y 8D dejaron auditoría técnica y sensibilidad descriptiva. El cierre continúa con la política predictiva v2/reubicación (8E), evaluación estadística final (8F), integración de resultados en pantalla (8G), instrumentos y evaluación con personas (fase 9), reproducibilidad de entrega y versión final del artículo."
    )

    document.add_heading("3. Datos: origen, ventana y preparación", level=1)
    document.add_heading("Fuente y periodo", level=2)
    document.add_paragraph(
        "La fuente es el dataset público ‘EMS Incident Dispatch Data’ de NYC Open Data, publicado en Socrata, recurso 76xm-jjuj. La extracción congelada del estudio contiene incidentes desde el 1 de abril hasta el 30 de junio de 2026 (fin exclusivo: 1 de julio), en hora de Nueva York. Se descargó el 8 de septiembre de 2026, hora de Lima, y se conservó como snapshot identificable para que futuras actualizaciones del portal no cambien silenciosamente los datos analizados."
    )
    add_table(document, ["Mes", "Incidentes"], [
        ["Abril 2026", "128,922"], ["Mayo 2026", "138,649"], ["Junio 2026", "138,427"], ["Total", "405,998"],
    ], widths=[3.3, 2.1])
    document.add_paragraph("Página del dataset: https://data.cityofnewyork.us/Public-Safety/EMS-Incident-Dispatch-Data/76xm-jjuj")
    document.add_paragraph("API de lectura: https://data.cityofnewyork.us/resource/76xm-jjuj.json")
    document.add_heading("Qué se hizo con los registros", level=2)
    for item in [
        "Se descargaron páginas ordenadas por fecha e identificador, se verificaron conteos e IDs y se guardaron localmente con manifiesto y hashes.",
        "Se perfiló calidad y se mantuvo una cohorte de demanda de 405,998 incidentes.",
        "La cohorte de respuesta contiene 365,578 registros con respuesta válida; 40,420 registros no cumplieron los criterios de respuesta y no se convierten en tiempos cero.",
        "La cohorte de transporte contiene 223,333 registros. Once casos con secuencias temporales incompatibles se trazaron en cuarentena.",
        "Para pronóstico, los incidentes se agregaron a una rejilla de borough-hora de cinco zonas y se generaron variables de calendario y rezagos históricos.",
    ]:
        add_bullet(document, item)
    add_callout(document, "Por qué guardar datos locales", "La API fue el medio de adquisición. El análisis usa un snapshot local versionado: eso permite repetir el estudio con la misma extracción, verificar integridad y no depender de que el portal siga exactamente igual.")
    document.add_heading("Qué información no trae la fuente", level=2)
    add_table(document, ["No está observado en el snapshot", "Qué implica"], [
        ["GPS o trayectoria real de cada ambulancia", "Las ubicaciones/unidades del mapa son representaciones del escenario, no recorridos reales."],
        ["Inventario y turnos reales de la flota", "Los tamaños de flota son escenarios asumidos para sensibilidad."],
        ["Red vial y tráfico en cada momento", "Los tiempos entre zonas son aproximaciones, no navegación GPS."],
        ["Capacidad, camas y demoras hospitalarias completas", "La liberación de recursos contiene supuestos calibrados/proxy."],
        ["Variables socioeconómicas individuales", "No se pueden inferir quintiles, ingresos, raza ni equidad social solo con este dataset."],
    ])

    document.add_heading("4. Modelos predictivos y sus métricas", level=1)
    document.add_heading("Qué predicen", level=2)
    document.add_paragraph(
        "Los modelos estiman el número de incidentes que podría ocurrir en la siguiente hora en cada uno de los cinco boroughs. No predicen directamente el tiempo de respuesta de una ambulancia, el diagnóstico clínico ni la ubicación GPS de la próxima llamada. La predicción puede alimentar la regla de despacho predictiva del simulador."
    )
    document.add_heading("Modelos comparados", level=2)
    add_table(document, ["Modelo", "Idea", "Resultado en prueba"], [
        ["Baseline temporal", "Promedia demanda histórica según zona, día de semana y hora, con respaldos más generales.", "MAE 6.035 · RMSE 8.349 · WAPE 15.70%"],
        ["Random Forest", "Combina muchos árboles y promedia sus predicciones; captura relaciones no lineales.", "MAE 5.557 · RMSE 7.501 · WAPE 14.45%"],
        ["Gradient Boosting", "Agrega árboles secuencialmente para corregir errores anteriores.", "MAE 5.638 · RMSE 7.725 · WAPE 14.66%"],
    ])
    document.add_paragraph("Datos de prueba: 3,600 celdas borough-hora de junio de 2026. Menor MAE, RMSE y WAPE suele indicar menor error para este problema.")
    document.add_heading("Cómo explicar cada métrica", level=2)
    add_table(document, ["Métrica", "Lectura para el público", "En este estudio"], [
        ["MAE", "Error absoluto promedio. Se expresa en llamadas por borough-hora.", "Random Forest: 5.557. En promedio, la estimación se desvía unas 5.6 llamadas por celda evaluada."],
        ["RMSE", "Similar al MAE, pero penaliza más los errores grandes.", "Random Forest: 7.501; baseline: 8.349."],
        ["WAPE", "Suma de errores absolutos dividida entre la demanda observada total; expresado en porcentaje.", "14.45% significa que el error absoluto agregado equivale a 14.45% de las llamadas observadas en esas celdas."],
        ["Sesgo", "Promedio de predicción menos valor observado. Negativo indica subestimación promedio.", "Random Forest: -0.319; Gradient Boosting: -0.864; baseline: -1.942."],
    ])
    add_callout(document, "No hay matriz de confusión", "La demanda es un conteo numérico y el modelo resuelve regresión. Una matriz de confusión se usa para categorías; aquí son más adecuadas las métricas de error y las gráficas de valores reales frente a predichos.")
    document.add_heading("¿Por qué se seleccionó Gradient Boosting si Random Forest ganó en prueba?", level=2)
    document.add_paragraph(
        "La regla fijada selecciona el modelo con menor MAE en validación, antes de usar el periodo de prueba. Gradient Boosting obtuvo MAE 5.674 frente a 5.681 de Random Forest en validación, una diferencia pequeña. En prueba final, Random Forest consiguió MAE 5.557 y RMSE 7.501, ligeramente mejores que Gradient Boosting. Se deben comunicar ambos hechos: selección según el criterio previo y desempeño observado en test. Cambiar el modelo elegido después de mirar el test convertiría ese test en parte de la selección y sesgaría la evaluación."
    )
    add_callout(document, "Respuesta breve", "Gradient Boosting fue seleccionado por la regla de validación; Random Forest fue el que obtuvo el menor error en el conjunto de prueba. La diferencia observada es descriptiva y no prueba por sí sola que sea superior en otros periodos.", fill="FFF3E8", accent=ORANGE)
    document.add_heading("Cómo se entrenó y qué significa la validación temporal", level=2)
    document.add_paragraph(
        "La separación conserva el orden del tiempo para evitar entrenar con el futuro: entrenamiento [1 de abril, 16 de mayo), validación [16 de mayo, 1 de junio) y prueba [1 de junio, 1 de julio) de 2026. Para las variables de historial se usan rezagos de 1, 24 y 168 horas y ventanas móviles; estas se calculan solo con observaciones anteriores al periodo pronosticado. Además, el informe CRISP-DM contiene una evaluación complementaria rolling-origin de cinco ventanas semanales: el origen avanza y cada evaluación usa periodos posteriores al entrenamiento de ese fold. No es una partición aleatoria k-fold. Sin embargo, junio ya se consultó en análisis posteriores; por eso, aunque fue el bloque de prueba definido originalmente, hoy debe presentarse como evidencia retrospectiva y exploratoria, no como confirmación independiente ciega."
    )
    document.add_paragraph(
        "La búsqueda de hiperparámetros fue ligera y exploratoria, con tres configuraciones candidatas para cada ensamble. El artefacto de entrenamiento principal documenta para Random Forest 180 árboles, profundidad máxima 18 y mínimo de 3 muestras por hoja; para HistGradientBoosting, 260 iteraciones, tasa de aprendizaje 0.06, 31 nodos hoja y regularización L2 de 0.05. No se trata de una optimización exhaustiva ni de una garantía de parámetros óptimos."
    )
    add_callout(
        document,
        "Nota de consistencia metodológica",
        "La búsqueda complementaria CRISP-DM y el entrenamiento principal de Fase 5 son artefactos distintos. En la tabla exploratoria, el candidato 4 de Gradient Boosting (180 iteraciones, tasa 0.04) obtiene MAE 5.527 en validación, menor que el candidato 5 (260 iteraciones, tasa 0.06; MAE 5.674) usado en el modelo principal que alimenta la app. Por eso, presenta el resultado de Fase 5 como la comparación y selección implementada, y reconoce que la alineación entre búsqueda complementaria y modelo final debe cerrarse antes de afirmar que se desplegó el mejor de todos los candidatos explorados.",
        fill="FFF3E8",
        accent=ORANGE,
    )
    document.add_heading("Lectura de una tabla o matriz de métricas", level=2)
    document.add_paragraph(
        "Cada fila representa un modelo y cada columna una métrica calculada sobre el mismo conjunto de evaluación. Para MAE, RMSE y WAPE, valores menores son preferibles; para sesgo se busca cercanía a cero, además de revisar su dirección. Un mapa de calor solo codifica magnitudes con color: hay que leer los números y confirmar la escala de cada columna, porque las métricas tienen unidades distintas."
    )

    document.add_heading("5. Cómo funciona y se interpreta la simulación", level=1)
    document.add_paragraph(
        "El motor de eventos discretos recibe llamadas, las coloca en espera cuando no hay recurso disponible, asigna ambulancias, simula el tiempo hasta escena, la atención y, según el escenario, traslado y liberación de la unidad. El usuario configura fuente de demanda, política, flota asumida, modelo, hora inicial, duración y semilla. La semilla permite repetir el mismo escenario aleatorio."
    )
    add_table(document, ["Entrada / política", "Qué representa"], [
        ["Histórica", "Usa las horas, zonas y atributos disponibles de llamadas observadas. La atención y asignación posteriores se vuelven contrafactuales."],
        ["Pronosticada", "Genera llegadas sintéticas a partir de pronósticos horarios; no son incidentes futuros observados."],
        ["Tradicional", "Elige la unidad con menor tiempo estimado a escena según la regla v1."],
        ["Predictiva v1", "Combina tiempo estimado con una penalización heurística ligada a demanda pronosticada y unidades libres en la zona de origen."],
    ])
    document.add_heading("Qué significa cada indicador de la pantalla", level=2)
    add_table(document, ["Indicador", "Cómo leerlo"], [
        ["Llamadas recibidas", "Incidentes incorporados a la corrida hasta el reloj mostrado."],
        ["En espera / activas", "En espera: aún no asignadas. Activas: asignadas o en servicio y todavía no liberadas, según el estado del motor."],
        ["Flota disponible", "Unidades libres de la flota configurada para el escenario. No es un censo de ambulancias reales."],
        ["Respuesta simulada", "Tiempo promedio de llamada a llegada a escena entre los casos que ya tienen esa llegada observada por el motor."],
        ["P90 simulado", "Percentil 90 de los tiempos simulados con respuesta conocida: 90% de esos casos tiene tiempo igual o menor."],
        ["Utilización temporal", "Fracción de tiempo de la flota que estuvo ocupada; depende de la ventana y de los supuestos del servicio."],
        ["Referencia observada", "Respuesta histórica válida conocida hasta ese corte. Puede tener una población distinta de los resultados simulados."],
        ["Completados / sin completar", "Completado significa servicio terminado y unidad liberada. Sin completar incluye estados aún en proceso al corte; no significa necesariamente que nunca hubo llegada a escena."],
    ])
    add_callout(document, "Corte de tiempo", "Al finalizar una corrida puede haber llamadas esperando, unidades en ruta o servicios activos. Por eso una media entre quienes ya llegaron a escena puede parecer baja mientras quedan casos sin respuesta. Hay que leer juntos los tiempos, pendientes, completitud y momento de medición.", fill="FFF3E8", accent=ORANGE)
    document.add_heading("Qué muestra el mapa 2D/3D", level=2)
    for item in [
        "Los polígonos representan los cinco boroughs; los colores distinguen zonas.",
        "Los indicadores de ambulancias representan unidades del escenario y sus estados simulados. No son posiciones recibidas desde GPS.",
        "Las barras por borough muestran la demanda prevista de la hora activa; el pronóstico cambia con el reloj de simulación.",
        "Los trazos de movimiento son esquemáticos. Los edificios y elementos 3D sirven para visualizar la zona, no describen infraestructura EMS medida.",
        "La escena es una lectura espacial de los resultados del motor; el mapa por sí mismo no demuestra que la política sea mejor.",
    ]:
        add_bullet(document, item)

    document.add_heading("6. Resultados de sensibilidad de la simulación", level=1)
    document.add_paragraph(
        "La fase 8D completó 220 pares (440 corridas): 22 escenarios, diez semillas por escenario, ventanas de desarrollo de tres días de mayo, medición de 24 horas, calentamiento base de seis horas y vaciado de hasta 24 horas. Se compararon flotas asumidas de 691, 531, 400 y 300 unidades, perturbaciones exploratorias de ±20% en algunos supuestos y demanda histórica frente a demanda generada. Estas cifras no son un inventario oficial de FDNY."
    )
    add_table(document, ["Flota asumida", "Política", "Respuesta media tras vaciado", "Sin respuesta al corte", "Utilización"], [
        ["300", "Tradicional", "97.379 min", "648.47", "0.809"],
        ["300", "Predictiva v1", "96.809 min", "645.50", "0.806"],
        ["400", "Tradicional", "21.860 min", "54.27", "0.621"],
        ["400", "Predictiva v1", "18.913 min", "39.27", "0.616"],
        ["531", "Tradicional", "12.236 min", "32.67", "0.441"],
        ["531", "Predictiva v1", "12.296 min", "32.67", "0.441"],
        ["691", "Tradicional", "12.166 min", "32.67", "0.338"],
        ["691", "Predictiva v1", "12.166 min", "32.67", "0.338"],
    ])
    document.add_paragraph(
        "‘Sin respuesta al corte’ cuenta llamadas sin llegada a escena al terminar las 24 horas medidas, incluyendo espera sin asignación y unidades todavía en ruta. Tras el periodo de vaciado, las 440 corridas lograron respuesta para los incidentes de la cohorte medida; esto permite calcular resultados posteriores al corte, pero no borra el retraso que tuvieron durante la ventana."
    )
    add_figure(document, "fleet_response.png", "Figura: tiempo medio según tamaño de flota asumido", "Al pasar de 300 a 400 unidades, la respuesta media cae de forma marcada en este escenario. Con 531 y 691, los valores se acercan. Las bandas son el mínimo y máximo observados entre días y semillas; no son intervalos de confianza. La flota es hipotética y solo se estudiaron tres días de desarrollo.")
    add_figure(document, "fleet_unanswered.png", "Figura: llamadas sin llegada a escena en el corte", "A 300 unidades queda una cola considerable al terminar el día. El vaciado posterior ayuda a distinguir retrasos pendientes de casos no resueltos dentro del límite, pero la operación durante las 24 horas sigue siendo relevante. No comparar solo medias de respuesta de quienes ya fueron atendidos.")
    add_figure(document, "borough_response.png", "Figura: tiempo de respuesta por borough con 300 unidades asumidas", "Permite ver que un promedio global puede ocultar diferencias geográficas. Las diferencias entre boroughs no miden equidad por nivel socioeconómico, raza o género; esas variables no están integradas.")
    document.add_heading("¿La política predictiva ganó?", level=2)
    document.add_paragraph(
        "No se puede sostener una victoria general. En el escenario de 400 unidades, la media descriptiva de Predictiva v1 fue menor; con 531 fue ligeramente mayor; con 691 coincidieron a la precisión reportada. Los resultados son de desarrollo, con tres días, y la fase 8D no calculó pruebas inferenciales. La fase 8C también encontró escenarios donde ambas políticas tomaban las mismas decisiones, especialmente cuando había muchas unidades libres cerca. La conclusión defendible es que el comportamiento depende del nivel de recursos y de los supuestos; se requiere completar fases estadísticas posteriores antes de afirmar superioridad."
    )
    add_callout(document, "Cómo presentar los números", "Son resultados descriptivos de simulación con datos históricos y recursos asumidos. No son tiempos medidos en operaciones reales, ni una promesa de reducción para una ciudad.", fill="FFF3E8", accent=ORANGE)

    document.add_heading("7. Utilidad y adaptación al contexto local", level=1)
    document.add_paragraph(
        "La utilidad local sería metodológica y de planificación: organizar registros, identificar patrones horarios/territoriales, visualizar concentración de demanda y comparar escenarios de asignación antes de pilotear cambios. El software y el flujo analítico pueden adaptarse; las métricas, parámetros y predicciones de Nueva York no se trasladan automáticamente a Trujillo o al Perú."
    )
    add_table(document, ["Para una adaptación defendible", "Dato local que se necesitaría"], [
        ["Estimar demanda por hora y zona", "Fecha/hora de llamada, tipo/categoría y ubicación geográfica con nivel de precisión autorizado."],
        ["Medir tiempos de respuesta", "Tiempos de recepción, despacho, salida, llegada a escena y reglas consistentes de calidad."],
        ["Modelar capacidad real", "Ambulancias activas por turno, bases, estados, tripulación y disponibilidad."],
        ["Representar viajes", "Red vial, distancias/tiempos por franja, tráfico y restricciones locales."],
        ["Representar traslado y hospitales", "Destino, tiempos de traslado, espera/entrega y disponibilidad de centros, si se cuenta con autorización."],
        ["Evaluar desempeño con usuarios", "Participación de personal EMS, tareas, instrumentos, revisión ética y protección de datos."],
    ])
    add_qa(document, "¿Entonces sí serviría en Trujillo?", "La arquitectura y el método pueden servir como punto de partida para analizar el sistema local. Para que las predicciones y la simulación representen Trujillo se tendría que conseguir información local, rehacer el entrenamiento y calibrar el motor con operadores. Sin esos datos solo sería una demostración de la plataforma, no evidencia sobre el servicio peruano.")
    add_qa(document, "¿Se necesita GPS?", "Para evaluar un despacho operativo y ubicación dinámica, GPS sería muy valioso, junto con datos de despacho y disponibilidad. No es indispensable para iniciar una investigación agregada por zonas y tiempos históricos; sí limita cuánto puede afirmarse sobre rutas, ubicación exacta y desempeño en tiempo real.")
    add_qa(document, "¿Cuál sería el aporte investigativo?", "Una plataforma reproducible que integra pipeline histórico, pronóstico por zonas y comparación controlada de políticas dentro de un simulador, junto con evidencia de sensibilidad a la disponibilidad de recursos. El aporte local fuerte requeriría calibración y evaluación con datos y actores del contexto local.")

    document.add_heading("8. Preguntas frecuentes del público", level=1)
    add_qa(document, "¿El sistema ya está conectado a ambulancias reales?", "No. Usa un snapshot histórico local y ejecuta escenarios. No recibe ubicación GPS, disponibilidad ni llamadas en vivo.")
    add_qa(document, "¿La predicción indica dónde estará la siguiente ambulancia?", "No. Predice el conteo de demanda de la siguiente hora para cada borough. La selección de ambulancia ocurre después dentro del simulador.")
    add_qa(document, "¿Los 405,998 registros son todas las llamadas de Nueva York?", "Son los registros devueltos por el dataset público dentro de la ventana abril-junio de 2026 y reconciliados con sus conteos de extracción. No certifican por sí solos la completitud de los registros internos de la agencia.")
    add_qa(document, "¿Por qué se eligieron tres meses?", "Fue una ventana acotada para manejar volumen y construir un snapshot reproducible. Permite evaluar un piloto temporal, pero no captura un año entero ni demuestra estacionalidad anual.")
    add_qa(document, "¿Qué significa WAPE?", "Es el error absoluto acumulado dividido por el total observado. En prueba, 14.45% para Random Forest indica que la suma de errores absolutos equivale al 14.45% del total de llamadas observadas en las celdas evaluadas. No significa que cada predicción falle exactamente 14.45%.")
    add_qa(document, "¿Qué métrica es la más importante?", "Depende de la decisión. MAE ofrece un error promedio fácil de explicar; RMSE hace más visibles errores grandes; WAPE permite comparar el error con el volumen total. Conviene leerlas juntas y revisar diferencias por borough.")
    add_qa(document, "¿Por qué no usar MAPE?", "La demanda por hora puede ser cero en algunas celdas; MAPE se vuelve indefinido o inestable al dividir por cero o valores pequeños. WAPE evita esa división fila por fila y usa el volumen total como denominador.")
    add_qa(document, "¿Por qué no hay matriz de confusión?", "No se clasifican llamadas en categorías predichas; se estima un número de llamadas. La matriz de confusión no describe el error de regresión de conteos.")
    add_qa(document, "¿La simulación pronostica la demanda de las próximas 12 horas?", "No necesariamente. El pronóstico disponible es horario y retrospectivo, emitido para cada hora con el historial anterior. Una corrida de 12 horas alimentada por pronósticos no equivale a emitir desde el inicio una predicción de 12 horas completa.")
    add_qa(document, "¿Los tiempos observados y simulados se pueden restar para decir cuánto mejoró?", "No sin controlar que representen la misma población y denominador. La referencia observada y las respuestas contrafactuales pueden incluir casos distintos; una comparación de políticas requiere entradas pareadas y reporte de pendientes/completitud.")
    add_qa(document, "¿Qué son las 300, 400, 531 y 691 ambulancias?", "Tamaños hipotéticos de flota usados para analizar sensibilidad. Se derivaron como proxies de capacidad del estudio; no son cifras oficiales del inventario EMS de NYC ni una recomendación para Trujillo.")
    add_qa(document, "¿El resultado prueba que poner más ambulancias siempre resuelve el problema?", "En estos escenarios, más recursos reducen cola y respuesta, pero son configuraciones simuladas y no incorporan todos los costos, turnos, bases ni restricciones. El resultado describe el modelo y los supuestos evaluados; no es una recomendación operativa directa.")
    add_qa(document, "¿Qué papel cumple el asistente con LangGraph?", "Consulta documentación local, métricas y estado de simulación para contestar preguntas sobre el proyecto. No es el motor predictivo ni un agente autónomo que controle ambulancias.")
    add_qa(document, "¿Se usaron encuestas o participaron operadores?", "Las fases analíticas descritas usan datos administrativos y simulación. La evaluación con expertos/usuarios mediante validez de contenido, alfa de Cronbach o KR-20 corresponde a instrumentos por desarrollar/aplicar; esos coeficientes no validan el pronóstico ni sustituyen una prueba operacional.")

    document.add_heading("9. Límites, ética y lenguaje responsable", level=1)
    for item in [
        "El estudio es retrospectivo y de simulación; no prueba impacto causal ni desempeño en producción.",
        "Los datos de NYC no representan el sistema EMS peruano. Cualquier transferencia requiere datos locales y evaluación propia.",
        "Los datos públicos están agregados para proteger ubicaciones; el prototipo genera posiciones aproximadas para visualización y no debe llamarlas GPS reales.",
        "Las flotas, la distribución de recursos y varios tiempos de servicio son supuestos o proxies. Deben aparecer con su procedencia y análisis de sensibilidad.",
        "El análisis por borough no equivale a equidad socioeconómica. No se integraron variables de ingreso, raza o género.",
        "La predicción puede ayudar a explorar escenarios, pero la priorización clínica y las decisiones de despacho necesitan validación, gobernanza y supervisión del personal responsable.",
    ]:
        add_bullet(document, item)
    add_table(document, ["Evita decir", "Di mejor"], [
        ["‘La IA reduce los tiempos de respuesta en NYC.’", "‘En escenarios simulados, observamos diferencias que dependen de la flota asumida; todavía no es una evaluación operacional.’"],
        ["‘El mapa muestra la ubicación real de las ambulancias.’", "‘El mapa presenta posiciones y estados del escenario simulado.’"],
        ["‘El modelo predice emergencias.’", "‘El modelo estima conteos horarios de demanda por borough.’"],
        ["‘La flota óptima es 400 ambulancias.’", "‘400 fue uno de los niveles hipotéticos comparados; no se optimizó ni se validó como recomendación real.’"],
    ])
    document.add_heading("Estado de las pruebas estadísticas y de instrumentos", level=2)
    document.add_paragraph(
        "El repositorio contiene análisis estadísticos exploratorios para errores de modelos y comparaciones de políticas. La auditoría 8C documenta limitaciones que deben corregirse antes de usar valores p e intervalos como evidencia inferencial final: dependencia temporal entre horas/zonas, tratamiento de empates y casos degenerados, alineación de pares y comparación múltiple. La sensibilidad 8D por sí misma reporta promedios descriptivos y no estima valores p ni declara superioridad. En el plan de cierre, la fase 8F está destinada a la evaluación estadística final."
    )
    document.add_paragraph(
        "V de Aiken, alfa de Cronbach y KR-20 corresponden a instrumentos con jueces o participantes: validez de contenido, consistencia de escalas Likert y consistencia de preguntas dicotómicas. No se calculan a partir de las llamadas EMS y no miden la exactitud de los modelos. Su aplicación requiere recoger valoraciones/respuestas reales; no debe afirmarse que ya se validó el sistema con esos coeficientes."
    )

    document.add_heading("10. Glosario para la exposición", level=1)
    add_table(document, ["Término", "Significado sencillo"], [
        ["Borough", "Una de las cinco divisiones territoriales de la ciudad de Nueva York."],
        ["Snapshot", "Copia local congelada de una versión de los datos, con fecha e identificador."],
        ["Cohorte", "Subconjunto de registros que cumple una regla analítica definida."],
        ["Baseline", "Método de referencia sencillo para saber si un modelo más complejo aporta mejora."],
        ["Validación", "Periodo usado para comparar/configurar candidatos antes de la evaluación final."],
        ["Prueba / test", "Periodo reservado para medir desempeño final; debe mantenerse fuera de la selección."],
        ["Rezago / lag", "Demanda observada en una hora anterior, usada como variable para pronosticar."],
        ["Contrafactual", "Resultado posible bajo una política alternativa aplicada a las mismas entradas simuladas."],
        ["P90", "Valor bajo el cual cae el 90% de las observaciones; muestra la cola de tiempos altos."],
        ["Calentamiento", "Tiempo previo a la ventana medida para que el estado inicial no empiece artificialmente vacío."],
        ["Vaciado", "Tiempo posterior al fin de llegadas que permite observar cómo se resuelven las llamadas pendientes."],
        ["Semilla", "Número que hace reproducibles los componentes aleatorios de una corrida."],
        ["Utilización", "Parte del tiempo disponible en que los recursos simulados permanecen ocupados."],
    ])

    document.add_heading("11. Dónde consultar evidencia del proyecto", level=1)
    add_table(document, ["Tema", "Documento o sección"], [
        ["Extracción y snapshot", "docs/NYC_EMS_FULL_EXTRACTION_REPORT.md"],
        ["Calidad y cohortes", "docs/NYC_EMS_PHASE3_QUALITY_REPORT.md"],
        ["Variables y particiones", "docs/NYC_EMS_PHASE4_FEATURES_REPORT.md"],
        ["Resultados de modelos", "docs/NYC_EMS_PHASE5_MODEL_REPORT.md y reports/nyc_ems/<snapshot>/model_performance/"],
        ["Calibración y supuestos", "docs/NYC_EMS_PHASE6_SIMULATION_CALIBRATION_REPORT.md"],
        ["Integración frontend/API", "docs/NYC_EMS_PHASE7_API_FRONTEND_REPORT.md"],
        ["Experimentos de políticas", "docs/NYC_EMS_PHASE8_EXPERIMENTS.md y docs/NYC_EMS_PHASE8C_AUDIT.md"],
        ["Sensibilidad 8D", "docs/NYC_EMS_PHASE8D_SENSITIVITY.md y reports/nyc_ems/<snapshot>/phase8d/583b13134c76f924dffd/"],
        ["Método CRISP-DM", "docs/NYC_EMS_CRISP_DM_METHODOLOGY.md"],
        ["Plan de implementación y cierre", "docs/PLAN_IMPLEMENTACION_NYC_EMS_2026.md y docs/PLAN_CIERRE_IMPLEMENTACION_NYC_EMS.md"],
    ])
    document.add_paragraph(
        "Snapshot usado: " + SNAPSHOT + ". Los datos crudos y muchos artefactos son locales y están excluidos de Git; para reproducir la misma evidencia se debe conservar la carpeta del snapshot y seguir los comandos de research/README.md."
    )
    add_callout(document, "Última idea para recordar", "Explique siempre tres cosas juntas: qué dato se observó, qué parte fue simulada y qué supuestos conectan ambas. Esa distinción hace que la demostración sea clara y creíble.")

    document.core_properties.title = "Guía de exposición: Gemelo Digital EMS"
    document.core_properties.subject = "Descripción del proyecto, métricas, simulación y preguntas frecuentes"
    document.core_properties.author = "Proyecto EMS Digital Twin"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
